import logging
import re
from pathlib import Path
from typing import Any, Dict, List

from app.services.chunking_service import chunking_service

logger = logging.getLogger("faceattend.document_ingestion")

BACKEND_ROOT = Path(__file__).resolve().parents[2]
DOCUMENTS_DIR = BACKEND_ROOT / "data" / "documents"
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB


class DocumentIngestionService:
    """Loads supported documents, extracts text, chunks it, and adds metadata."""

    def __init__(self, documents_dir: Path = DOCUMENTS_DIR):
        self.documents_dir = documents_dir
        self.documents_dir.mkdir(parents=True, exist_ok=True)

    @staticmethod
    def _safe_filename(filename: str) -> str:
        filename = Path(filename).name
        filename = re.sub(r"[^A-Za-z0-9._-]+", "_", filename)
        return filename[:180] or "uploaded_document"

    def save_upload(self, filename: str, content: bytes) -> Path:
        if not filename:
            raise ValueError("A filename is required.")

        suffix = Path(filename).suffix.lower()
        if suffix not in ALLOWED_EXTENSIONS:
            raise ValueError(
                f"Unsupported file type '{suffix}'. Allowed: {', '.join(sorted(ALLOWED_EXTENSIONS))}"
            )

        if len(content) > MAX_FILE_SIZE:
            raise ValueError("File is too large. Maximum allowed size is 10 MB.")

        safe_name = self._safe_filename(filename)
        destination = self.documents_dir / safe_name
        destination.write_bytes(content)
        return destination

    def _extract_txt_md(self, path: Path) -> List[Dict[str, Any]]:
        text = path.read_text(encoding="utf-8", errors="ignore")
        return [{"text": text, "page": None}]

    def _extract_pdf(self, path: Path) -> List[Dict[str, Any]]:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError(
                "PDF ingestion requires pypdf. Install it with: pip install pypdf"
            ) from exc

        reader = PdfReader(str(path))
        pages = []
        for page_number, page in enumerate(reader.pages, start=1):
            pages.append({
                "text": page.extract_text() or "",
                "page": page_number,
            })
        return pages

    def _extract_docx(self, path: Path) -> List[Dict[str, Any]]:
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError(
                "DOCX ingestion requires python-docx. Install it with: pip install python-docx"
            ) from exc

        document = Document(str(path))
        paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]

        # Include table content because policies are often represented in tables.
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    paragraphs.append(" | ".join(cells))

        return [{"text": "\n".join(paragraphs), "page": None}]

    def extract_pages(self, path: Path) -> List[Dict[str, Any]]:
        suffix = path.suffix.lower()
        if suffix in {".txt", ".md"}:
            return self._extract_txt_md(path)
        if suffix == ".pdf":
            return self._extract_pdf(path)
        if suffix == ".docx":
            return self._extract_docx(path)
        raise ValueError(f"Unsupported file type: {suffix}")

    def ingest_file(self, path: Path) -> List[Dict[str, Any]]:
        pages = self.extract_pages(path)
        all_chunks: List[Dict[str, Any]] = []

        for page_data in pages:
            page_chunks = chunking_service.chunk_text(
                page_data["text"],
                metadata={
                    "source": path.name,
                    "source_path": str(path),
                    "document_type": "uploaded_document",
                    "file_type": path.suffix.lower().lstrip("."),
                    "page": page_data.get("page"),
                },
            )
            all_chunks.extend(page_chunks)

        results = []
        for chunk in all_chunks:
            results.append({
                "id": f"upload_{path.stem}_{chunk['metadata']['chunk_index']}",
                "title": f"{path.name} — Chunk {chunk['metadata']['chunk_index'] + 1}",
                "category": "Uploaded Document",
                "content": chunk["content"],
                "metadata": chunk["metadata"],
            })

        logger.info("Ingested %s: %s chunks", path.name, len(results))
        return results

    def ingest_all(self) -> List[Dict[str, Any]]:
        documents: List[Dict[str, Any]] = []
        for path in sorted(self.documents_dir.iterdir()):
            if path.is_file() and path.suffix.lower() in ALLOWED_EXTENSIONS:
                try:
                    documents.extend(self.ingest_file(path))
                except Exception:
                    logger.exception("Failed to ingest document: %s", path.name)
        return documents

    def list_documents(self) -> List[Dict[str, Any]]:
        return [
            {
                "filename": p.name,
                "size_bytes": p.stat().st_size,
                "extension": p.suffix.lower(),
            }
            for p in sorted(self.documents_dir.iterdir())
            if p.is_file() and p.suffix.lower() in ALLOWED_EXTENSIONS
        ]


document_ingestion_service = DocumentIngestionService()
